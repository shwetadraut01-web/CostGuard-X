import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def markdown_to_docx(md_filepath: str, docx_filepath: str, doc_title: str):
    doc = docx.Document()
    
    # Page setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = title_p.add_run(doc_title)
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42) # slate-900

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("CostGuard-X: Cloud Cost Waste Intelligence Platform (No-ML Engine)")
    sub_run.font.name = 'Arial'
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(71, 85, 105) # slate-600

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    with open(md_filepath, "r", encoding="utf-8") as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    in_table = False
    table_data = []

    for line in lines:
        raw_line = line
        line = line.strip()

        # Code block handler
        if line.startswith("```"):
            if in_code_block:
                # End code block
                p_code = doc.add_paragraph()
                p_code.paragraph_format.left_indent = Inches(0.3)
                p_code.paragraph_format.space_before = Pt(4)
                p_code.paragraph_format.space_after = Pt(8)
                run_c = p_code.add_run("\n".join(code_lines))
                run_c.font.name = 'Consolas'
                run_c.font.size = Pt(9.5)
                run_c.font.color.rgb = RGBColor(30, 41, 59)
                in_code_block = False
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(raw_line.rstrip())
            continue

        # Table handler
        if line.startswith("|") and line.endswith("|"):
            if "---" in line:
                continue
            cells = [c.strip() for c in line.split("|")[1:-1]]
            table_data.append(cells)
            in_table = True
            continue
        elif in_table:
            # End of table
            if table_data:
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True

                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = val
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        run = p.runs[0] if p.runs else p.add_run()
                        run.font.name = 'Arial'
                        run.font.size = Pt(9.5)

                        if r_idx == 0:
                            set_cell_background(cell, "1E293B") # Dark slate
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                        else:
                            bg_color = "F1F5F9" if r_idx % 2 == 1 else "FFFFFF"
                            set_cell_background(cell, bg_color)
                            run.font.color.rgb = RGBColor(15, 23, 42)

                doc.add_paragraph().paragraph_format.space_after = Pt(8)
            in_table = False
            table_data = []

        if not line:
            continue

        # Headings
        if line.startswith("# "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            r = h.add_run(line[2:].replace("**", "").replace("*", ""))
            r.font.name = 'Arial'
            r.font.size = Pt(18)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith("## "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            r = h.add_run(line[3:].replace("**", "").replace("*", ""))
            r.font.name = 'Arial'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 58, 138) # blue-900
        elif line.startswith("### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            r = h.add_run(line[4:].replace("**", "").replace("*", ""))
            r.font.name = 'Arial'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            text = line[2:]
            r = p.add_run(text.replace("**", "").replace("*", ""))
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(51, 65, 85)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line[2:].replace("**", ""))
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.italic = True
            r.font.color.rgb = RGBColor(71, 85, 105)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(line.replace("**", "").replace("*", ""))
            r.font.name = 'Arial'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(30, 41, 59)

    doc.save(docx_filepath)
    print(f"Successfully generated: {docx_filepath}")

if __name__ == "__main__":
    desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop")
    
    # 1. Export Master Document
    master_md = "docs/CostGuard_X_Project_Master_Document.md"
    master_docx = os.path.join(desktop_dir, "CostGuard_X_Project_Master_Document.docx")
    if os.path.exists(master_md):
        markdown_to_docx(master_md, master_docx, "CostGuard-X Master Project Documentation")

    # 2. Export Research Paper Specification
    paper_md = "docs/research_paper_specification.md"
    if not os.path.exists(paper_md):
        paper_md = r"C:\Users\hp\.gemini\antigravity\brain\ac9edb07-cd0d-4916-aaa0-885a81f4878a\research_paper_specification.md"
    paper_docx = os.path.join(desktop_dir, "CostGuard_X_Research_Paper_Specification.docx")
    if os.path.exists(paper_md):
        markdown_to_docx(paper_md, paper_docx, "CostGuard-X Academic Research Paper Specification")
